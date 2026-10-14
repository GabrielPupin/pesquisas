"""Tabular views of a result (shared by CSV, Excel and the PDF report)."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from ..analysis import AnalysisResult
from ..model.project import Project


def _fmt_dist(t):
    return t.distribution.describe() + f" [{t.distribution_origin}]" if t.distribution is not None else \
        (t.distribution_origin or "none")


def tables(project: Project, res: AnalysisResult) -> dict[str, tuple[list[str], list[list]]]:
    u = project.base_unit
    out = {}
    out["Inputs"] = (
        ["id", "name", "component", "kind", "sensitivity", "sign origin", f"nominal [{u}]", f"upper [{u}]",
         f"lower [{u}]", "as specified", "distribution (deviation)", "source"],
        [[t.contributor_id, t.name, t.component, t.kind, t.sensitivity, t.sign_origin, t.nominal, t.upper, t.lower,
          t.tolerance_text, _fmt_dist(t), t.source_text] for t in res.diagnosis.terms])
    used = {project.contributor(t.contributor_id).geometric_id for t in res.diagnosis.terms
            if t.kind == "geometric"}
    out["Geometric"] = (
        ["id", "type", "feature", "zone", "datums", "modifier", "standard", "size principle", "propagated",
         "effect distribution", "notes"],
        [[g.id, g.type, g.feature_id, ("Ø" if g.diameter_zone else "") + f"{g.value:g} {g.unit}",
          " | ".join(g.datums), g.material_modifier or "", g.standard or "", g.size_principle or "undeclared",
          "yes" if g.id in used else "no", g.effect_distribution.describe() if g.effect_distribution else "",
          g.notes] for g in project.geometric_tolerances])
    out["Checks"] = (["id", "category", "title", "outcome", "severity", "details", "action"],
                     [[c.id, c.category, c.title, c.outcome, c.severity or "", " ; ".join(c.details), c.action]
                      for c in res.diagnosis.checks])
    out["Sensitivity"] = (
        ["rank", "id", "name", "sensitivity", f"tolerance width [{u}]", "WC share", f"effect on min [{u}]",
         f"effect on max [{u}]", f"σ [{u}]", "variance share"],
        [[r.rank, r.contributor_id, r.name, r.sensitivity, r.tolerance_width, r.wc_share, r.effect_on_min,
          r.effect_on_max, r.sigma, r.variance_share] for r in res.sensitivity])
    rows = []
    if res.worst_case:
        w = res.worst_case
        rows += [["Worst-Case", "nominal", w.nominal], ["Worst-Case", "minimum", w.minimum],
                 ["Worst-Case", "maximum", w.maximum], ["Worst-Case", "exact", w.exact]]
    if res.rss:
        r = res.rss
        rows += [["RSS", "mean", r.mean], ["RSS", "sigma", r.sigma], ["RSS", f"mean-{r.k:g}σ", r.lower_k],
                 ["RSS", f"mean+{r.k:g}σ", r.upper_k], ["RSS", "P(y<LSL)", r.p_below], ["RSS", "P(y>USL)", r.p_above]]
    if res.monte_carlo:
        m = res.monte_carlo
        rows += [["Monte Carlo", "n", m.n], ["Monte Carlo", "seed", m.seed], ["Monte Carlo", "mean", m.mean],
                 ["Monte Carlo", "std", m.std], ["Monte Carlo", "median", m.median],
                 ["Monte Carlo", "observed min (not WC)", m.observed_min],
                 ["Monte Carlo", "observed max (not WC)", m.observed_max]]
        rows += [["Monte Carlo", f"P{p:g}", v] for p, v in m.percentiles.items()]
        rows += [["Monte Carlo", "P(y<LSL)", m.p_below], ["Monte Carlo", "P(y>USL)", m.p_above],
                 ["Monte Carlo", "P(nonconforming)", m.p_total],
                 ["Monte Carlo", f"CI{m.confidence * 100:.0f}% low", m.ci_total[0] if m.ci_total else None],
                 ["Monte Carlo", f"CI{m.confidence * 100:.0f}% high", m.ci_total[1] if m.ci_total else None],
                 ["Monte Carlo", "invalid samples", m.invalid_samples]]
    for k, v in res.skipped.items():
        rows.append([k.upper(), "not computed", v])
    out["Results"] = (["method", "quantity", f"value [{u}] / probability"], rows)
    if res.monte_carlo:
        out["MC convergence"] = (["n", "mean", "std", "p_fail"],
                                 [[c["n"], c["mean"], c["std"], c["p_fail"]] for c in res.monte_carlo.convergence])
    trace = []
    for t in res.diagnosis.terms:
        c = project.contributor(t.contributor_id)
        if c.source and c.source.callout_id:
            rec = next((d for d in project.drawings if any(x.id == c.source.callout_id for x in d.callouts)), None)
            conf = rec.confirmation_for(c.source.callout_id) if rec else None
            trace.append([t.contributor_id, c.source.document, c.source.revision, c.source.page, c.source.callout_id,
                          c.source.raw_text, conf.reviewer if conf else "", conf.timestamp if conf else "",
                          "yes" if conf and conf.edited else "no", rec.sha256[:16] if rec else ""])
        else:
            trace.append([t.contributor_id, c.source.document if c.source else "", c.source.revision if c.source else "",
                          "", "", "manual entry", "", "", "", ""])
    out["Traceability"] = (["contributor", "document", "revision", "page", "callout", "text on drawing", "confirmed by",
                            "confirmed at", "edited", "drawing sha256"], trace)
    out["Assumptions"] = (["#", "assumption"], [[i + 1, a] for i, a in enumerate(res.diagnosis.assumptions)])
    out["Missing information"] = (["#", "item"], [[i + 1, a] for i, a in enumerate(res.diagnosis.missing_information)])
    out["Meta"] = (["key", "value"], [[k, str(v)] for k, v in res.meta.items()] +
                   [["status", res.diagnosis.status], ["equation", res.diagnosis.equation]] +
                   [["conclusion", c] for c in res.conclusions])
    return out


def export_csv(project: Project, res: AnalysisResult, directory: str | Path) -> list[Path]:
    d = Path(directory)
    d.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, (head, rows) in tables(project, res).items():
        p = d / (name.lower().replace(" ", "_") + ".csv")
        with p.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(head)
            w.writerows(rows)
        paths.append(p)
    return paths


def export_excel(project: Project, res: AnalysisResult, path: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, (head, rows) in tables(project, res).items():
        ws = wb.create_sheet(name[:31])
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([x if isinstance(x, (int, float, str)) or x is None else str(x) for x in r])
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    p = Path(path)
    wb.save(p)
    return p
