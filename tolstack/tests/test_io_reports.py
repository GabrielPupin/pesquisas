"""Persistence, reports and CLI."""

import json

import pytest
from openpyxl import load_workbook

from tolstack.analysis import run_analysis
from tolstack.cli import main
from tolstack.io.project_json import load_project, project_from_dict, project_to_json, save_project
from tolstack.errors import ModelError
from tolstack.reporting.pdf import export_pdf
from tolstack.reporting.tabular import export_csv, export_excel


def test_json_roundtrip_preserves_fingerprint(axial_project, tmp_path):
    f = tmp_path / "p.json"
    save_project(axial_project, f)
    q = load_project(f)
    assert q.fingerprint() == axial_project.fingerprint()
    data = json.loads(project_to_json(axial_project))
    data["schema_version"] = "2.0"
    with pytest.raises(ModelError):
        project_from_dict(data)


def test_duplicate_ids_rejected(axial_project):
    data = json.loads(project_to_json(axial_project))
    data["contributors"].append(data["contributors"][0])
    with pytest.raises(ModelError):
        project_from_dict(data)


def test_reports_are_written(fit_project, tmp_path):
    res = run_analysis(fit_project, mc_samples=5_000)
    pdf = export_pdf(fit_project, res, tmp_path / "r.pdf")
    assert pdf.read_bytes()[:4] == b"%PDF" and pdf.stat().st_size > 5_000
    x = export_excel(fit_project, res, tmp_path / "r.xlsx")
    wb = load_workbook(x)
    assert {"Inputs", "Checks", "Results", "Sensitivity", "Traceability", "Geometric"} <= set(wb.sheetnames)
    paths = export_csv(fit_project, res, tmp_path / "csv")
    assert any(p.name == "results.csv" for p in paths)


def test_result_traceability(axial_project):
    res = run_analysis(axial_project, "A", mc_samples=5_000, seed=7)
    assert res.meta["fingerprint"] == axial_project.fingerprint()
    assert res.meta["seed"] == 7 and res.meta["mc_samples"] == 5_000
    assert res.warnings and "not been approved" in res.warnings[0]
    d = res.as_dict()
    assert d["diagnosis"]["status"] == "Conditionally Valid"
    json.dumps(d, default=str)


def test_cli(tmp_path, demos, capsys):
    f = tmp_path / "ax.json"
    save_project(demos.axial_demo(), f)
    assert main(["check", str(f)]) == 0
    assert main(["analyze", str(f), "--chain", "B"]) == 2
    assert main(["analyze", str(f), "--chain", "A", "--n", "2000", "--xlsx", str(tmp_path / "a.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "Worst-Case" in out and "Sensitivity ranking" in out
